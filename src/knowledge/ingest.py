"""Ingestion of AFS material into separated evidence records.

Each public ``ingest_*`` function returns a list of :class:`Evidence` records
with a stable ``chunk_id`` and an explicit :class:`SourceKind`, so that primary
PDF text, verified S02 notes, the spreadsheet transcription, the development
reference and business context stay strictly separated.

Production dependencies (PyMuPDF as ``fitz``, ``python-docx``, ``openpyxl``,
PyYAML) are imported lazily inside the functions that need them. The sandbox
used to author this ticket has PyPI blocked, so those libraries cannot be
installed or run here; the real index build runs locally on the maintainer's
machine. The pure-logic helpers (chunk-id formatting, note parsing, allowlist)
are exercised by the test-suite without any of those libraries.

Allowlist
---------
``INDEXABLE_SOURCES`` is the hard allowlist of source_ids that may enter the
technical evidence index. It deliberately EXCLUDES:
    * customer_scenarios_for_final_evaluation.txt (H01-H06, final batch input)
    * the QA-schedule PDFs (organisational, no technical evidence)
    * the development reference DOCX (stored separately as dev_reference)
"""

from __future__ import annotations

import os
import re
from typing import Any, Optional

from . import catalog
from .types import Evidence, SourceKind

# ---------------------------------------------------------------------------
# Allowlist: which sources may enter data/index/evidence.jsonl.
# DOCX (dev_reference) is handled separately and is intentionally NOT here.
# ---------------------------------------------------------------------------
INDEXABLE_SOURCES: frozenset[str] = frozenset(
    {"BRIEF", "S01", "S02", "S03", "S04", "S05", "S06", "XLSX"}
)

# PDF sources ingested page-by-page as plain text (kind=pdf_text).
PDF_TEXT_SOURCES: dict[str, str] = {
    "BRIEF": "AI-based molecular media analysis by process-spectroscopy in-line.pdf",
    "S01": "S01_Multispectral_Laboratory.pdf",
    "S03": "S03_LED_Kidney_Study_2017.pdf",
    "S04": "S04_Combi_MIR_Fluorescence_2021.pdf",
}

# Files that must NEVER be indexed, regardless of anything else.
EXCLUDED_FROM_INDEX: frozenset[str] = frozenset(
    {
        "customer_scenarios_for_final_evaluation.txt",
        "Spectroscopy_Assistant_QA_schedule_17_October.pdf",
        "Spectroscopy_Assistant_QA_schedule_17_October (1).pdf",
    }
)

S06_URL = "https://mirphotonics.com/laboratory/"
S02_NOTES_DEFAULT = os.path.join("data", "source_notes", "s02_notes.yaml")


# ---------------------------------------------------------------------------
# Stable chunk-id helpers (deterministic, no hashing needed).
# ---------------------------------------------------------------------------
def pdf_chunk_id(source_id: str, page: int) -> str:
    """e.g. ``S01-p010`` (1-based, zero-padded to 3 digits)."""

    return f"{source_id}-p{page:03d}"


def table_chunk_id(row_index: int) -> str:
    """e.g. ``XLSX-row03`` (1-based data-row index, zero-padded to 2 digits)."""

    return f"XLSX-row{row_index:02d}"


def table_note_chunk_id(sheet: str) -> str:
    """e.g. ``XLSX-notes-notes`` for the provenance ``notes`` sheet."""

    return f"XLSX-notes-{sheet}"


def s02_note_chunk_id(page: int) -> str:
    """e.g. ``S02-note-p13``."""

    return f"S02-note-p{page}"


def business_chunk_id(source_id: str) -> str:
    """e.g. ``S05-context`` / ``S06-context``."""

    return f"{source_id}-context"


# ---------------------------------------------------------------------------
# AP3.1 - PDF page text (PyMuPDF). Used for S01, S03, S04, BRIEF.
# ---------------------------------------------------------------------------
def ingest_pdf_text(path: str | os.PathLike[str], source_id: str) -> list[Evidence]:
    """Ingest one evidence record per PDF page (1-based), kind=pdf_text.

    Requires PyMuPDF (imported as ``fitz``). Raises ``ModuleNotFoundError`` if
    the library is unavailable, so the caller can report the environment gap.
    """

    import fitz  # PyMuPDF; lazily imported (not installable in sandbox)

    records: list[Evidence] = []
    with fitz.open(os.fspath(path)) as doc:
        for zero_based, pymupdf_page in enumerate(doc):
            page = zero_based + 1  # 1-based, counting the title page
            text = pymupdf_page.get_text("text").strip()
            if not text:
                # Image-only page: keep a placeholder so the page still exists
                # in the index, but do not invent content.
                text = f"[no extractable text on {source_id} p.{page}]"
            records.append(
                Evidence(
                    source_id=source_id,
                    page=page,
                    text=text,
                    kind=SourceKind.pdf_text,
                    chunk_id=pdf_chunk_id(source_id, page),
                )
            )
    return records


# ---------------------------------------------------------------------------
# AP3.2 - Verified S02 slide notes (hand-written YAML).
# ---------------------------------------------------------------------------
def _parse_s02_notes_yaml(raw: str) -> dict[str, Any]:
    """Minimal parser for the specific s02_notes.yaml structure.

    Production uses PyYAML (``yaml.safe_load``); this fallback understands only
    the subset of YAML the notes file uses: a ``meta`` mapping and a ``notes``
    list of ``page/title/text`` entries with folded (``>-``) scalars. It lets
    the tests run without PyYAML installed.
    """

    try:
        import yaml  # type: ignore

        return yaml.safe_load(raw) or {}
    except ModuleNotFoundError:
        pass

    data: dict[str, Any] = {"meta": {}, "notes": []}
    lines = raw.splitlines()
    i = 0
    n = len(lines)

    def strip_comment(s: str) -> str:
        return s

    # --- meta block ---
    while i < n:
        line = lines[i]
        if line.strip().startswith("#") or not line.strip():
            i += 1
            continue
        if line.rstrip() == "meta:":
            i += 1
            while i < n and (lines[i].startswith("  ") and ":" in lines[i]) and not lines[i].lstrip().startswith("#"):
                key, _, val = lines[i].strip().partition(":")
                val = val.strip().strip('"')
                if val in ("true", "false"):
                    data["meta"][key] = val == "true"
                else:
                    data["meta"][key] = val
                i += 1
            continue
        if line.rstrip() == "notes:":
            i += 1
            break
        i += 1

    # --- notes list ---
    current: Optional[dict[str, Any]] = None
    while i < n:
        raw_line = lines[i]
        stripped = raw_line.strip()
        if stripped.startswith("#"):
            i += 1
            continue
        m = re.match(r"^\s*-\s*page:\s*(\d+)\s*$", raw_line)
        if m:
            if current:
                data["notes"].append(current)
            current = {"page": int(m.group(1))}
            i += 1
            continue
        if current is not None:
            tm = re.match(r'^\s*title:\s*"?(.*?)"?\s*$', raw_line)
            if tm:
                current["title"] = tm.group(1)
                i += 1
                continue
            if re.match(r"^\s*text:\s*>-?\s*$", raw_line):
                # Folded scalar: gather deeper-indented lines.
                base_indent = len(raw_line) - len(raw_line.lstrip())
                i += 1
                parts: list[str] = []
                while i < n:
                    nxt = lines[i]
                    if not nxt.strip():
                        i += 1
                        continue
                    indent = len(nxt) - len(nxt.lstrip())
                    if indent <= base_indent:
                        break
                    parts.append(nxt.strip())
                    i += 1
                current["text"] = " ".join(parts)
                continue
        i += 1
    if current:
        data["notes"].append(current)
    return data


def ingest_verified_s02_notes(
    notes_path: str | os.PathLike[str] = S02_NOTES_DEFAULT,
) -> list[Evidence]:
    """Load hand-verified S02 slide notes, kind=verified_note.

    Graceful on a missing file: returns an empty list.
    """

    notes_path = os.fspath(notes_path)
    if not os.path.isfile(notes_path):
        return []

    with open(notes_path, "r", encoding="utf-8") as fh:
        data = _parse_s02_notes_yaml(fh.read())

    meta = data.get("meta", {}) or {}
    provenance = meta.get(
        "provenance",
        "reviewed transcription from AFS source catalogue; pending final visual check",
    )

    records: list[Evidence] = []
    for note in data.get("notes", []) or []:
        page = note.get("page")
        title = (note.get("title") or "").strip()
        body = (note.get("text") or "").strip()
        text = f"{title}. {body}" if title else body
        text = f"{text} [provenance: {provenance}]"
        records.append(
            Evidence(
                source_id="S02",
                page=page,
                text=text,
                kind=SourceKind.verified_note,
                chunk_id=s02_note_chunk_id(page),
            )
        )
    return records


# ---------------------------------------------------------------------------
# AP3.3 - Spreadsheet transcription (openpyxl). S01 Table 1 (p.6).
# ---------------------------------------------------------------------------
def _format_table_row(headers: list[str], values: list[Any]) -> str:
    pairs = []
    for h, v in zip(headers, values):
        if v is None or str(v).strip() == "":
            continue
        pairs.append(f"{h}={v}")
    return "; ".join(pairs)


def ingest_table_xlsx(path: str | os.PathLike[str]) -> list[Evidence]:
    """Ingest the method-comparison spreadsheet, kind=table_transcription.

    Emits one record per data row of ``method_comparison`` plus records for the
    ``notes`` sheet. All records cite ``S01 Table 1 (p.6)`` and flag the
    sensitivity/selectivity/cost ratings as qualitative, never as measured
    specifications. Requires openpyxl.
    """

    from openpyxl import load_workbook  # lazily imported (not in sandbox)

    QUALIFIER = (
        " [transcription of S01 Table 1 (p.6); sensitivity/selectivity/cost are "
        "qualitative ratings (low/average/high), not measured specifications]"
    )

    wb = load_workbook(os.fspath(path), read_only=True, data_only=True)
    records: list[Evidence] = []

    if "method_comparison" in wb.sheetnames:
        ws = wb["method_comparison"]
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            data_index = 0
            for row in rows[1:]:
                if row is None or all(c is None for c in row):
                    continue
                data_index += 1
                text = _format_table_row(headers, list(row)) + QUALIFIER
                records.append(
                    Evidence(
                        source_id="XLSX",
                        page=6,
                        text=text,
                        kind=SourceKind.table_transcription,
                        chunk_id=table_chunk_id(data_index),
                    )
                )

    if "notes" in wb.sheetnames:
        ws = wb["notes"]
        note_lines = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                note_lines.append(" ".join(cells))
        if note_lines:
            text = " ".join(note_lines) + QUALIFIER
            records.append(
                Evidence(
                    source_id="XLSX",
                    page=6,
                    text=text,
                    kind=SourceKind.table_transcription,
                    chunk_id=table_note_chunk_id("notes"),
                )
            )

    wb.close()
    return records


# ---------------------------------------------------------------------------
# AP3.4 - Development reference (python-docx). D01-D08 + worked answers.
# Stored SEPARATELY as dev_reference; never part of evidence.jsonl.
# ---------------------------------------------------------------------------
def ingest_dev_reference_docx(path: str | os.PathLike[str]) -> list[Evidence]:
    """Extract the development reference DOCX, kind=dev_reference.

    Returns paragraph-level records tagged dev_reference so later tickets (T05
    few-shot) can use them. These are kept out of the technical evidence index.
    Requires python-docx.
    """

    import docx  # python-docx; lazily imported (not installable in sandbox)

    document = docx.Document(os.fspath(path))
    records: list[Evidence] = []
    index = 0
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        index += 1
        records.append(
            Evidence(
                source_id="DOCX",
                page=None,
                text=text,
                kind=SourceKind.dev_reference,
                chunk_id=f"DOCX-p{index:04d}",
            )
        )
    return records


# ---------------------------------------------------------------------------
# AP3.5 - Business context: S05 flyer + S06 link.
# ---------------------------------------------------------------------------
def ingest_business_context(
    materials_dir: str | os.PathLike[str],
) -> list[Evidence]:
    """Ingest S05 (flyer) and S06 (link) as kind=business_context.

    S06 is a URL, not a PDF page, so its evidence has ``page=None``. S05 is an
    image-heavy flyer: only its catalog description is recorded, no invented
    content. Graceful on missing files.
    """

    materials_dir = os.fspath(materials_dir)
    records: list[Evidence] = []

    # S05 - flyer (image-heavy). Use the verified catalog description only.
    s05 = catalog.get_entry("S05")
    s05_path = os.path.join(materials_dir, s05.filename) if s05 else None
    if s05 and s05_path and os.path.isfile(s05_path):
        records.append(
            Evidence(
                source_id="S05",
                page=1,
                text=(
                    f"{s05.description} [business context only; image-heavy flyer, "
                    "not evidence of technical performance]"
                ),
                kind=SourceKind.business_context,
                chunk_id=business_chunk_id("S05"),
            )
        )

    # S06 - laboratory link (URL only, not a PDF page => page=None).
    s06 = catalog.get_entry("S06")
    s06_path = os.path.join(materials_dir, s06.filename) if s06 else None
    url = S06_URL
    if s06_path and os.path.isfile(s06_path):
        try:
            with open(s06_path, "r", encoding="utf-8") as fh:
                content = fh.read().strip()
            if content:
                url = content.splitlines()[0].strip()
        except OSError:
            pass
    if s06:
        records.append(
            Evidence(
                source_id="S06",
                page=None,
                text=f"AFS laboratory website: {url} [business context only]",
                kind=SourceKind.business_context,
                chunk_id=business_chunk_id("S06"),
            )
        )

    return records


__all__ = [
    "INDEXABLE_SOURCES",
    "PDF_TEXT_SOURCES",
    "EXCLUDED_FROM_INDEX",
    "S06_URL",
    "pdf_chunk_id",
    "table_chunk_id",
    "table_note_chunk_id",
    "s02_note_chunk_id",
    "business_chunk_id",
    "ingest_pdf_text",
    "ingest_verified_s02_notes",
    "ingest_table_xlsx",
    "ingest_dev_reference_docx",
    "ingest_business_context",
]
